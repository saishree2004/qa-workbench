import io
import os
import re

from docx import Document
from docx.shared import Inches

from db import get_conn

FIELDS = [("Severity", "severity"), ("Description", "description"), ("Steps to Reproduce", "steps"),
          ("Expected Result", "expected"), ("Actual Result", "actual"), ("Suggested Fix", "suggested_fix")]


def bugs_to_docx(bug_ids) -> bytes:
    doc = Document()
    fig = 0
    with get_conn() as c:
        for bid in bug_ids:
            b = c.execute("SELECT * FROM bugs WHERE id=?", (bid,)).fetchone()
            if not b:
                continue
            doc.add_heading(b["title"] or "Untitled", level=1)
            for label, key in FIELDS:
                p = doc.add_paragraph()
                p.add_run(label + ":").bold = True
                text = b[key] or ""
                if key == "steps":
                    lines = [re.sub(r"^\s*\d+[.)]\s*", "", s).strip() for s in text.splitlines()]
                    for n, s in enumerate([s for s in lines if s], 1):
                        doc.add_paragraph(f"{n}. {s}")
                elif key == "severity":
                    p.add_run(" " + text)
                else:
                    doc.add_paragraph(text)
            for img in c.execute("SELECT * FROM bug_images WHERE bug_id=? ORDER BY fig_no, id", (bid,)):
                if not os.path.isfile(img["path"] or ""):
                    continue
                fig += 1
                doc.add_picture(img["path"], width=Inches(6))
                doc.add_paragraph(f"Fig. {fig}: {img['caption'] or ''}")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
